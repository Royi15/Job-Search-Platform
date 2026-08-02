"""Plain-text extraction for non-PDF document sources (PDF lives in
ats_parser.py — reused as-is, not duplicated here)."""
from pptx import Presentation


def extract_pptx_text(path: str) -> str:
    prs = Presentation(path)
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if text.strip():
                    lines.append(text)
    return "\n".join(lines)
